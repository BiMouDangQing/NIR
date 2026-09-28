"""把「秋月梨光谱数据含RealValue-解密.xlsx」转换为训练用 CSV。

源格式（宽表，多列块拼接）：
    [RealValue, Pixel, WaveL, DarkR, DarkS_S, DarkS_M, DarkS_L, REF,
     Sample(1), Sample(2), ..., Sample(K)] × 16 个块
其中每块：RealValue 列有 K 个非空值（写在前 K 行），块内有 K 个 Sample 列，
每个 Sample 列是完整 1024 维光谱（1024 行）。

目标格式（长表，两列）：
    RealValue,sample
    <标签>,"v1,v2,...,v1024"

即每行 = 一个样本：标签 + 逗号分隔的 1024 维光谱；无关字段（Pixel/WaveL/DarkR/.../REF）舍弃。
"""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC = PROJECT_ROOT / "data" / "秋月梨光谱数据含RealValue-解密.xlsx"
DST = PROJECT_ROOT / "data" / "秋月梨光谱数据_RealValue_sample.csv"


def convert(src: Path = SRC, dst: Path = DST) -> tuple[int, int]:
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    ncol = len(header)

    rv_idx = [
        i for i, h in enumerate(header) if isinstance(h, str) and h.startswith("RealValue")
    ]
    rv_idx.append(ncol)  # 哨兵，便于取最后一块

    total = 0
    with open(dst, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["RealValue", "sample"])
        for b in range(len(rv_idx) - 1):
            start, end = rv_idx[b], rv_idx[b + 1]
            sample_cols = [
                i
                for i in range(start + 1, end)
                if isinstance(header[i], str) and header[i].startswith("Sample")
            ]
            block = list(
                ws.iter_rows(min_row=2, min_col=start + 1, max_col=end, values_only=True)
            )
            n_rows = len(block)
            # 标签：RealValue 列（块内第 0 列）按行序的非空值
            labels = [row[0] for row in block if row[0] is not None]
            # 第 j 个标签 ↔ 第 j 个 Sample 列
            for j, sc in enumerate(sample_cols):
                if j >= len(labels):
                    break
                rel = sc - start
                spec = [block[r][rel] for r in range(n_rows)]
                spec_str = ",".join(str(v) for v in spec)
                writer.writerow([labels[j], spec_str])
                total += 1
    wb.close()
    return total, len(rv_idx) - 1


if __name__ == "__main__":
    total, blocks = convert()
    print(f"转换完成：{blocks} 个块，共 {total} 个样本")
    print(f"输出: {DST}")
